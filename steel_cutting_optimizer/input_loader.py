"""Đọc và validate dữ liệu đầu vào.

Hỗ trợ 2 nguồn:
1. File Excel `input.xlsx` với 2 sheet "Stock" + "Required" (khuyến nghị).
2. 2 file CSV riêng (stock.csv + required.csv) — giữ tương thích ngược.

Validate: length > 0, quantity > 0, không trùng id.
"""

from __future__ import annotations

import csv
from pathlib import Path

from models import Component, Stock


class InputValidationError(ValueError):
    """Lỗi khi dữ liệu đầu vào không hợp lệ."""


# ============================================================
# Nguồn Excel (input.xlsx) — khuyến nghị
# ============================================================


def load_from_excel(excel_path: str | Path) -> tuple[list[Stock], list[Component]]:
    """Đọc input.xlsx với 2 sheet "Stock" và "Required".

    Raises:
        InputValidationError: nếu thiếu sheet, thiếu cột, hoặc dữ liệu không hợp lệ.
    """
    from openpyxl import load_workbook

    path = Path(excel_path)
    if not path.exists():
        raise InputValidationError(f"Không tìm thấy file input Excel: {path}")

    wb = load_workbook(path, data_only=True)
    stocks: list[Stock] = []
    components: list[Component] = []

    # ---- Sheet Stock ----
    if "Stock" not in wb.sheetnames:
        raise InputValidationError(
            f"input.xlsx thiếu sheet 'Stock'. Có: {wb.sheetnames}"
        )
    ws_stock = wb["Stock"]
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws_stock[1]]
    required_stock_cols = ["stock_type_id", "length_mm", "quantity"]
    if not set(required_stock_cols).issubset(set(headers)):
        raise InputValidationError(
            f"Sheet 'Stock' thiếu cột. Cần: {required_stock_cols}, có: {headers}"
        )
    idx = {h: i for i, h in enumerate(headers)}
    seen_stock: set[str] = set()
    for row_no, row in enumerate(
        ws_stock.iter_rows(min_row=2, values_only=True), start=2
    ):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        sid = (
            str(row[idx["stock_type_id"]]).strip()
            if row[idx["stock_type_id"]] is not None
            else ""
        )
        try:
            length = float(row[idx["length_mm"]])
            qty = int(row[idx["quantity"]])
        except (TypeError, ValueError) as e:
            raise InputValidationError(
                f"Sheet 'Stock' dòng {row_no}: length_mm/quantity phải là số."
            ) from e
        if not sid:
            raise InputValidationError(
                f"Sheet 'Stock' dòng {row_no}: stock_type_id trống."
            )
        if sid in seen_stock:
            raise InputValidationError(
                f"Sheet 'Stock' dòng {row_no}: trùng stock_type_id '{sid}'."
            )
        if length <= 0:
            raise InputValidationError(
                f"Sheet 'Stock' dòng {row_no}: length_mm phải > 0."
            )
        if qty <= 0:
            raise InputValidationError(
                f"Sheet 'Stock' dòng {row_no}: quantity phải > 0."
            )
        seen_stock.add(sid)
        stocks.append(Stock(stock_type_id=sid, length_mm=length, quantity=qty))

    # ---- Sheet Required ----
    if "Required" not in wb.sheetnames:
        raise InputValidationError(
            f"input.xlsx thiếu sheet 'Required'. Có: {wb.sheetnames}"
        )
    ws_req = wb["Required"]
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws_req[1]]
    required_comp_cols = ["component_id", "length_mm", "quantity"]
    if not set(required_comp_cols).issubset(set(headers)):
        raise InputValidationError(
            f"Sheet 'Required' thiếu cột. Cần: {required_comp_cols}, có: {headers}"
        )
    idx = {h: i for i, h in enumerate(headers)}
    seen_comp: set[str] = set()
    for row_no, row in enumerate(
        ws_req.iter_rows(min_row=2, values_only=True), start=2
    ):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        cid = (
            str(row[idx["component_id"]]).strip()
            if row[idx["component_id"]] is not None
            else ""
        )
        try:
            length = float(row[idx["length_mm"]])
            qty = int(row[idx["quantity"]])
        except (TypeError, ValueError) as e:
            raise InputValidationError(
                f"Sheet 'Required' dòng {row_no}: length_mm/quantity phải là số."
            ) from e
        if not cid:
            raise InputValidationError(
                f"Sheet 'Required' dòng {row_no}: component_id trống."
            )
        if cid in seen_comp:
            raise InputValidationError(
                f"Sheet 'Required' dòng {row_no}: trùng component_id '{cid}'."
            )
        if length <= 0:
            raise InputValidationError(
                f"Sheet 'Required' dòng {row_no}: length_mm phải > 0."
            )
        if qty <= 0:
            raise InputValidationError(
                f"Sheet 'Required' dòng {row_no}: quantity phải > 0."
            )
        seen_comp.add(cid)
        components.append(Component(component_id=cid, length_mm=length, quantity=qty))

    wb.close()
    return stocks, components


# ============================================================
# Nguồn CSV (stock.csv + required.csv) — tương thích ngược
# ============================================================


def load_stock(csv_path: str | Path) -> list[Stock]:
    """Đọc stock.csv và validate."""
    path = Path(csv_path)
    if not path.exists():
        raise InputValidationError(f"Không tìm thấy file stock: {path}")

    stocks: list[Stock] = []
    seen_ids: set[str] = set()

    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        required_cols = {"stock_type_id", "length_mm", "quantity"}
        if reader.fieldnames is None or not required_cols.issubset(
            required_cols.intersection(set(reader.fieldnames))
        ):
            raise InputValidationError(
                f"stock.csv thiếu cột. Cần: {sorted(required_cols)}, có: {reader.fieldnames}"
            )

        for line_no, row in enumerate(reader, start=2):
            sid = (row.get("stock_type_id") or "").strip()
            try:
                length = float(row["length_mm"])
                qty = int(row["quantity"])
            except (TypeError, ValueError) as e:
                raise InputValidationError(
                    f"stock.csv dòng {line_no}: length_mm/quantity phải là số."
                ) from e

            if not sid:
                raise InputValidationError(
                    f"stock.csv dòng {line_no}: stock_type_id trống."
                )
            if sid in seen_ids:
                raise InputValidationError(
                    f"stock.csv dòng {line_no}: trùng stock_type_id '{sid}'."
                )
            if length <= 0:
                raise InputValidationError(
                    f"stock.csv dòng {line_no}: length_mm phải > 0."
                )
            if qty <= 0:
                raise InputValidationError(
                    f"stock.csv dòng {line_no}: quantity phải > 0."
                )

            seen_ids.add(sid)
            stocks.append(Stock(stock_type_id=sid, length_mm=length, quantity=qty))

    return stocks


def load_required(csv_path: str | Path) -> list[Component]:
    """Đọc required.csv và validate."""
    path = Path(csv_path)
    if not path.exists():
        raise InputValidationError(f"Không tìm thấy file required: {path}")

    components: list[Component] = []
    seen_ids: set[str] = set()

    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        required_cols = {"component_id", "length_mm", "quantity"}
        if reader.fieldnames is None or not required_cols.issubset(
            required_cols.intersection(set(reader.fieldnames))
        ):
            raise InputValidationError(
                f"required.csv thiếu cột. Cần: {sorted(required_cols)}, có: {reader.fieldnames}"
            )

        for line_no, row in enumerate(reader, start=2):
            cid = (row.get("component_id") or "").strip()
            try:
                length = float(row["length_mm"])
                qty = int(row["quantity"])
            except (TypeError, ValueError) as e:
                raise InputValidationError(
                    f"required.csv dòng {line_no}: length_mm/quantity phải là số."
                ) from e

            if not cid:
                raise InputValidationError(
                    f"required.csv dòng {line_no}: component_id trống."
                )
            if cid in seen_ids:
                raise InputValidationError(
                    f"required.csv dòng {line_no}: trùng component_id '{cid}'."
                )
            if length <= 0:
                raise InputValidationError(
                    f"required.csv dòng {line_no}: length_mm phải > 0."
                )
            if qty <= 0:
                raise InputValidationError(
                    f"required.csv dòng {line_no}: quantity phải > 0."
                )

            seen_ids.add(cid)
            components.append(
                Component(component_id=cid, length_mm=length, quantity=qty)
            )

    return components
