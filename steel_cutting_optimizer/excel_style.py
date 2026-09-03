"""Định dạng Excel dùng chung — giúp bảng dễ nhìn, dễ theo dõi.

Áp dụng cho cả input.xlsx và cutting_plan.xlsx.
- Header: nền màu + chữ trắng in đậm + viền.
- Banded rows: tô màu xen kẽ dòng.
- Freeze pane + auto filter trên header.
- Căn lề + độ rộng cột tự động.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# ===== Màu sắc =====
HEADER_FILL = PatternFill("solid", fgColor="2F5597")  # xanh dương đậm
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
BANDED_FILL = PatternFill("solid", fgColor="DDEBF7")  # xanh nhạt xen kẽ
SCRAP_FILL = PatternFill("solid", fgColor="FCE4EC")  # hồng nhạt (phế liệu)
WARN_FONT = Font(color="C00000", bold=True)  # đỏ cảnh báo
THIN_SIDE = Side(style="thin", color="999999")
BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")


def style_header(ws: Worksheet, row: int = 1) -> None:
    """Tô màu header: nền đậm + chữ trắng đậm + viền + canh giữa."""
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def apply_banded_rows(ws: Worksheet, start_row: int = 2) -> None:
    """Tô màu xen kẽ dòng (banded rows) từ start_row."""
    for row_idx in range(start_row, ws.max_row + 1):
        if (row_idx - start_row) % 2 == 0:
            for cell in ws[row_idx]:
                if cell.fill.fgColor.rgb in ("00000000", None):
                    cell.fill = BANDED_FILL


def add_borders(ws: Worksheet, start_row: int = 1) -> None:
    """Thêm viền cho toàn bộ vùng có dữ liệu."""
    for row in ws.iter_rows(
        min_row=start_row, min_col=1, max_row=ws.max_row, max_col=ws.max_column
    ):
        for cell in row:
            if cell.value is not None:
                cell.border = BORDER


def auto_width(ws: Worksheet, min_width: int = 10, max_width: int = 40) -> None:
    """Tự chỉnh độ rộng cột theo nội dung dài nhất."""
    for col_cells in ws.columns:
        max_len = 0
        letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value is not None:
                # Ưu tiên độ dài thực, tính cả tiếng Việt (len(str))
                length = len(str(cell.value))
                if length > max_len:
                    max_len = length
        width = max(min_width, min(max_len + 3, max_width))
        ws.column_dimensions[letter].width = width


def freeze_and_filter(ws: Worksheet, freeze: str = "A2") -> None:
    """Đóng băng dòng header + bật auto filter."""
    ws.freeze_panes = freeze
    if ws.max_row >= 1:
        ws.auto_filter.ref = ws.dimensions
