"""Tạo data/input.xlsx với 2 sheet Stock + Required, có định dạng đẹp.

Dữ liệu mẫu lấy từ dữ liệu thực tế đang có (S6000 + fc344...fc93).
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

import excel_style as es

BASE = Path(__file__).resolve().parent / "data"
OUT = BASE / "input.xlsx"

STOCK_ROWS = [
    ("S6000", 6000, 500),
]

REQUIRED_ROWS = [
    ("fc344", 2066, 2),
    ("fc83", 7065, 3),
    ("fc65", 7307, 4),
    ("fc61", 7302, 4),
    ("fc79", 7094, 4),
    ("fc66", 7340, 4),
    ("fc93", 7065, 1),
]


def main() -> None:
    wb = Workbook()

    # ---- Sheet Stock ----
    ws_stock = wb.active
    ws_stock.title = "Stock"
    ws_stock.append(["stock_type_id", "length_mm", "quantity"])
    for row in STOCK_ROWS:
        ws_stock.append(list(row))
    es.style_header(ws_stock)
    es.apply_banded_rows(ws_stock)
    es.add_borders(ws_stock)
    es.freeze_and_filter(ws_stock, "A2")
    es.auto_width(ws_stock)
    # Số căn giữa
    for row in ws_stock.iter_rows(min_row=2, min_col=2, max_col=3):
        for cell in row:
            cell.alignment = es.CENTER

    # ---- Sheet Required ----
    ws_req = wb.create_sheet("Required")
    ws_req.append(["component_id", "length_mm", "quantity"])
    for row in REQUIRED_ROWS:
        ws_req.append(list(row))
    es.style_header(ws_req)
    es.apply_banded_rows(ws_req)
    es.add_borders(ws_req)
    es.freeze_and_filter(ws_req, "A2")
    es.auto_width(ws_req)
    for row in ws_req.iter_rows(min_row=2, min_col=2, max_col=3):
        for cell in row:
            cell.alignment = es.CENTER

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Đã tạo: {OUT}")


if __name__ == "__main__":
    main()
